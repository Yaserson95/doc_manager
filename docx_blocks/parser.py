# docx_parser.py
import re
import base64
import io
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.document import Document as _DocxDocument
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from .tags import find_blocks
from .table import DocxTableParser
from .formatter import DocxFormatter
from .format import FormatSpec


class DocxBlockParser:
    """
    Парсер текста с блоками {block ...}...{/block} и встроенными
    тегами {text ...}...{/text}, {image ...}...{/image}.
    На выходе — объект docx.Document.
    """

    def __init__(self):
        self.table_parser = DocxTableParser(self)
        self.formatter = DocxFormatter()

    # ------------------------------------------------------------------
    # Публичный метод
    # ------------------------------------------------------------------
    def parse(self, input_text: str) -> Document:
        doc = Document()
        for block_str in find_blocks(input_text):
            self._add_block_to_container(doc, block_str)
        return doc

    # ------------------------------------------------------------------
    # Приватные методы — парсинг параметров
    # ------------------------------------------------------------------
    def _parse_params(self, params_str: str) -> dict:
        """Извлекает пары ключ:"значение" из строки параметров."""
        pattern = r'(\w+):"([^"]*)"'
        matches = re.findall(pattern, params_str)
        return dict(matches)

    def _parse_block_type(self, type_str: str):
        """Разбирает тип блока. Возвращает (тип, уровень)."""
        if type_str is None:
            return ('p', None)
        if type_str.startswith('h:'):
            level = int(type_str.split(':')[1])
            return ('h', level)
        elif type_str.startswith('ul:'):
            level = int(type_str.split(':')[1])
            return ('ul', level)
        elif type_str.startswith('ol:'):
            level = int(type_str.split(':')[1])
            return ('ol', level)
        elif type_str == 'table':
            return ('table', None)
        else:
            return ('p', None)
            
    def _add_block_to_container(self, container, block_str):
        block_data = self._parse_block(block_str)
        if not block_data:
            return None
        type_info = block_data['type_info']
        block_format = block_data['block_format']
        content = block_data['content']
        params = block_data['params']

        p_type, level = type_info

        if p_type == 'table':
            self.table_parser.add_table(container, params, content)
            return None

        paragraph = self._add_paragraph_with_type(
            container, type_info, level)
        self._apply_block_format(paragraph, block_format)
        self._process_block_content(paragraph, content, block_format)
        return paragraph

    def _parse_text_style(self, style_str: str) -> dict:
        """Разбирает строку стилей текста, например 'bold,underline'."""
        styles = style_str.split(',')
        result = {}
        for style in styles:
            style = style.strip().lower()
            if style in ['bold', 'italic', 'underline']:
                result[style] = True
        return result

    def _parse_block_format(self, format_str: str) -> FormatSpec:
        """Разбирает строку форматирования блока."""
        return FormatSpec.parse(format_str)
        
    def _parse_block(self, block_str: str) -> dict | None:
        """Разбирает один блок, возвращает словарь с данными."""
        m = re.match(r'\{block\s*(.*?)\}(.*)\{\/block\}\s*$',
                     block_str, re.DOTALL)
        if not m:
            return None
        params_str = m.group(1)
        content = m.group(2)
        params = self._parse_params(params_str)
        type_str = params.get('type', 'p')
        format_str = params.get('format', '')
        type_info = self._parse_block_type(type_str)
        block_format = self._parse_block_format(format_str)
        return {
            'params': params,
            'type_info': type_info,
            'block_format': block_format,
            'content': content
        }

    # ------------------------------------------------------------------
    # Приватные методы — работа с docx
    # ------------------------------------------------------------------
    def _parse_size(self, value):
        return self.formatter.parse_size(value)

    def _set_paragraph_shading(self, paragraph, color: str):
        """Устанавливает цвет фона для абзаца."""
        pPr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), color.lstrip('#'))
        pPr.append(shd)

    def _apply_line_height(self, paragraph, value: str):
        """Применяет высоту строки к абзацу."""
        value = value.strip().lower()
        if value.endswith(('pt', 'px', 'in', 'cm')):
            size = self._parse_size(value)
            if size:
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                paragraph.paragraph_format.line_spacing = size
        else:
            try:
                multiple = float(value)
                paragraph.paragraph_format.line_spacing = multiple
            except ValueError:
                pass

    def _apply_block_format(self, paragraph, spec):
        self.formatter.apply_paragraph(paragraph, spec)
    
    def _apply_run_format(self, run, block_spec, text_spec=None):
        spec = block_spec if block_spec is not None else FormatSpec()
        if text_spec is not None:
            spec = spec.merge(text_spec)
        self.formatter.apply_run(run, spec)
    
    def _add_paragraph_with_type(self, container, type_info, level):
        p_type, _ = type_info

        if isinstance(container, _Cell):
            p_el = OxmlElement('w:p')
            container._tc.append(p_el)
            p = Paragraph(p_el, container)
        else:
            p = container.add_paragraph()

        if p_type == 'h':
            p.style = f'Heading {level or 1}'
        elif p_type == 'ul':
            p.style = 'List Bullet'
            if level and level > 1:
                p.paragraph_format.left_indent = Inches(0.25 * (level - 1))
        elif p_type == 'ol':
            p.style = 'List Number'
            if level and level > 1:
                p.paragraph_format.left_indent = Inches(0.25 * (level - 1))
        return p

    # ------------------------------------------------------------------
    # Приватные методы — разбор содержимого блока
    # ------------------------------------------------------------------
    def _parse_content(self, content: str) -> list:
        """Разбирает содержимое блока на текстовые фрагменты и встроенные теги."""
        pattern = r'(\{text\s+.*?\}.*?\{\/text\}|\{image\s+.*?\}.*?\{\/image\})'
        parts = re.split(pattern, content, flags=re.DOTALL)
        result = []
        for i, part in enumerate(parts):
            if i % 2 == 0:
                if part:
                    result.append(('text', part))
            else:
                if part.startswith('{text'):
                    m = re.match(r'\{text\s+(.*?)\}(.*?)\{\/text\}', part, re.DOTALL)
                    if m:
                        result.append(('text_tag', m.group(1), m.group(2)))
                elif part.startswith('{image'):
                    m = re.match(r'\{image\s+(.*?)\}(.*?)\{\/image\}', part, re.DOTALL)
                    if m:
                        result.append(('image_tag', m.group(1), m.group(2)))
        return result

    def _add_image_to_paragraph(self, paragraph, image_params_str: str, image_content: str):
        """Добавляет изображение в абзац."""
        params = self._parse_params(image_params_str)
        alt_text = params.get('text', '')
        try:
            if image_content.startswith('data:image'):
                header, data = image_content.split(',', 1)
                image_data = base64.b64decode(data)
                image_stream = io.BytesIO(image_data)
                run = paragraph.add_run()
                run.add_picture(image_stream, width=Inches(6))
            else:
                run = paragraph.add_run()
                run.add_picture(image_content, width=Inches(6))
        except Exception:
            run = paragraph.add_run(f"[Image: {alt_text}]")

    def _process_block_content(self, paragraph, content: str, block_format: dict):
        """Обрабатывает содержимое блока и добавляет его в абзац."""
        items = self._parse_content(content)
        for item in items:
            if item[0] == 'text':
                text = item[1]
                if text:
                    run = paragraph.add_run(text)
                    self._apply_run_format(run, block_format)
            elif item[0] == 'text_tag':
                params_str = item[1]
                inner_text = item[2]
                params = self._parse_params(params_str)
                text_format_str = params.get('format', '')
                text_format = self._parse_block_format(text_format_str)
                run = paragraph.add_run(inner_text)
                self._apply_run_format(run, block_format, text_format)
            elif item[0] == 'image_tag':
                params_str = item[1]
                image_content = item[2]
                self._add_image_to_paragraph(paragraph, params_str, image_content)