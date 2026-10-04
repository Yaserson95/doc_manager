"""
Поддержка таблиц в блочном формате:
парсер DocxTableParser и сериализатор DocxTableSerializer.
"""

import re

from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.oxml.table import CT_Tbl
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

from .format import FormatSpec
from .tags import find_rows, find_cells, find_blocks

_BORDER_STYLE_MAP = {
        'solid':   'single',
        'dashed':  'dashed',
        'dotted':  'dotted',
        'double':  'double',
        'none':    'none',
        'hidden':  'none',
        # OOXML-стили, которые уже корректны
        'single':  'single',
        'dashdot': 'dashDot',
        'dashdotdot': 'dashDotDot',
        'dashsmallgap': 'dashSmallGap',
        'dotdash': 'dotDash',
        'dotdotdash': 'dotDotDash',
        'thick':   'thick',
        'thickthinlargegap': 'thickThinLargeGap',
        'thickthinsmallgap': 'thickThinSmallGap',
        'thickthinmediumgap': 'thickThinMediumGap',
        'thinlargegap': 'thinLargeGap',
        'thinmediumgap': 'thinMediumGap',
        'thinsmallgap': 'thinSmallGap',
        'thinThickLargeGap': 'thinThickLargeGap',
        'wave':    'wave',
        'triple':  'triple',
    }

# ======================================================================
#                                ПАРСЕР
# ======================================================================
class _CellData:
    __slots__ = ('params', 'blocks', 'rowspan', 'colspan', 'row', 'col')

    def __init__(self):
        self.params = {}
        self.blocks = []
        self.rowspan = 1
        self.colspan = 1
        self.row = 0
        self.col = 0


class DocxTableParser:
    """Парсит {block type:"table"}...{/block} и добавляет таблицу в контейнер.
    Для вложенных блоков использует DocxBlockParser (передаётся ссылкой).
    """

    def __init__(self, block_parser):
        self.bp = block_parser

    # ------------------------------------------------------------------
    # Публичный метод
    # ------------------------------------------------------------------
    def add_table(self, container, table_params, table_content):
        rows_data = self._parse_rows(table_content)
        if not rows_data:
            return None

        cells, num_cols, num_rows = self._layout_grid(rows_data)
        table = self._create_table(container, num_rows, num_cols)

        # --- формат таблицы ---
        table_spec = FormatSpec.parse(table_params.get('format', ''))
        self.bp.formatter.apply_table(table, table_spec)

        # --- объединения ---
        for cd in cells:
            if cd.rowspan > 1 or cd.colspan > 1:
                tl = table.cell(cd.row, cd.col)
                br = table.cell(cd.row + cd.rowspan - 1,
                                cd.col + cd.colspan - 1)
                if tl is not br:
                    tl.merge(br)

        # --- наполнение ячеек ---
        for cd in cells:
            self._fill_cell(table, cd)

        # --- формат строк ---
        for r_idx, row_data in enumerate(rows_data):
            row_spec = FormatSpec()
            for k, v in row_data['params'].items():
                row_spec.set(k, v)
            self.bp.formatter.apply_row(
                table.rows[r_idx], row_spec)

        return table

    # ------------------------------------------------------------------
    # Разбор
    # ------------------------------------------------------------------
    def _parse_rows(self, content):
        rows = []
        for row_str in find_rows(content):
            m = re.match(r'\{row\s*(.*?)\}(.*)\{/row\}\s*$',
                         row_str, re.DOTALL)
            if not m:
                continue
            row_params = self.bp._parse_params(m.group(1))
            row_inner = m.group(2)

            cells = []
            for cell_str in find_cells(row_inner):
                cm = re.match(r'\{cell\s*(.*?)\}(.*)\{/cell\}\s*$',
                              cell_str, re.DOTALL)
                if not cm:
                    continue
                cell_params = self.bp._parse_params(cm.group(1))
                cell_inner = cm.group(2)
                block_strs = find_blocks(cell_inner)
                cells.append({'params': cell_params, 'blocks': block_strs})

            rows.append({'params': row_params, 'cells': cells})
        return rows

    def _layout_grid(self, rows_data):
        """Жадно раскладывает ячейки в сетку с учётом colspan/rowspan."""
        max_cols = 0
        for row in rows_data:
            total = 0
            for c in row['cells']:
                try:
                    total += int(c['params'].get('colspan', 1))
                except (ValueError, TypeError):
                    total += 1
            max_cols = max(max_cols, total)

        num_rows = len(rows_data)
        occupancy = [[None] * max_cols for _ in range(num_rows)]
        result = []

        for r_idx, row in enumerate(rows_data):
            col = 0
            for cell in row['cells']:
                while col < max_cols and occupancy[r_idx][col] is not None:
                    col += 1
                if col >= max_cols:
                    break
                try:
                    cs = int(cell['params'].get('colspan', 1))
                except (ValueError, TypeError):
                    cs = 1
                try:
                    rs = int(cell['params'].get('rowspan', 1))
                except (ValueError, TypeError):
                    rs = 1
                cs = max(1, min(cs, max_cols - col))
                rs = max(1, min(rs, num_rows - r_idx))

                cd = _CellData()
                cd.params = cell['params']
                cd.blocks = cell['blocks']
                cd.colspan = cs
                cd.rowspan = rs
                cd.row = r_idx
                cd.col = col

                for rr in range(r_idx, r_idx + rs):
                    for cc in range(col, col + cs):
                        occupancy[rr][cc] = cd
                result.append(cd)
                col += cs

        return result, max_cols, num_rows

    # ------------------------------------------------------------------
    # Создание / наполнение
    # ------------------------------------------------------------------
    def _create_table(self, container, rows, cols):
        """Создаёт таблицу в Document или в _Cell."""
        if hasattr(container, '_tc'):
            # _Cell
            tbl_el = CT_Tbl.new_tbl(rows, cols, Inches(6))
            container._tc.append(tbl_el)
            table = DocxTable(tbl_el, container)
        else:
            table = container.add_table(rows=rows, cols=cols)
        try:
            table.style = 'Table Grid'
        except KeyError:
            pass
        return table

    def _fill_cell(self, table, cd):
        cell = table.cell(cd.row, cd.col)

        # Убираем дефолтные w:p
        tc = cell._tc
        for p in tc.findall(qn('w:p')):
            tc.remove(p)

        # Спецификация: top-level params + format:"..."
        cell_spec = FormatSpec()
        for k, v in cd.params.items():
            if k in ('colspan', 'rowspan', 'format'):
                continue
            cell_spec.set(k, v)
        cell_spec = cell_spec.merge(
            FormatSpec.parse(cd.params.get('format', '')))

        # Оформление ячейки (bg-color, valign, width, padding)
        self.bp.formatter.apply_cell(cell, cell_spec)

        # Значения по умолчанию для вложенных блоков.
        # align тоже прокидываем сюда — это его единственный рабочий путь.
        nested_default = FormatSpec()
        for k in ('text-style', 'txt-color', 'font-size',
                  'line-height', 'align'):
            if k in cell_spec:
                nested_default.set(k, cell_spec.args(k))

        for block_str in cd.blocks:
            bd = self.bp._parse_block(block_str)
            if not bd:
                continue
            type_info = bd['type_info']
            block_spec = bd['block_format']
            content = bd['content']

            if type_info[0] == 'table':
                self.bp.table_parser.add_table(
                    cell, bd['params'], content)
                continue

            paragraph = self.bp._add_paragraph_with_type(
                cell, type_info, type_info[1])
            merged = nested_default.merge(block_spec)
            self.bp.formatter.apply_paragraph(paragraph, merged)
            self.bp._process_block_content(paragraph, content, merged)

        if not cell.paragraphs:
            cell.add_paragraph()
    
    def _apply_cell_format(self, cell, spec: FormatSpec):
        tcPr = cell._tc.get_or_add_tcPr()

        # bg-color
        if 'bg-color' in spec:
            for old in tcPr.findall(qn('w:shd')):
                tcPr.remove(old)
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), str(spec.arg('bg-color')).lstrip('#'))
            tcPr.append(shd)

        # valign
        if 'valign' in spec:
            v = {'top': 'top', 'middle': 'center',
                 'bottom': 'bottom'}.get(str(spec.arg('valign')).lower())
            if v:
                for old in tcPr.findall(qn('w:vAlign')):
                    tcPr.remove(old)
                va = OxmlElement('w:vAlign')
                va.set(qn('w:val'), v)
                tcPr.append(va)

        # align
        if 'align' in spec:
            a = {'left': 'left', 'center': 'center',
                 'right': 'right', 'justify': 'both'}.get(
                     str(spec.arg('align')).lower())
            if a:
                for old in tcPr.findall(qn('w:jc')):
                    tcPr.remove(old)
                jc = OxmlElement('w:jc')
                jc.set(qn('w:val'), a)
                tcPr.append(jc)

        # width
        if 'width' in spec:
            size = self.bp._parse_size(spec.arg('width'))
            if size is not None:
                cell.width = size

        # padding
        self._apply_cell_padding(tcPr, spec)

    def _apply_cell_padding(self, tcPr, spec: FormatSpec):
        paddings = spec.padding()
        if not paddings:
            return
        for old in tcPr.findall(qn('w:tcMar')):
            tcPr.remove(old)
        mar = OxmlElement('w:tcMar')
        for side, tag in (('left', 'w:left'), ('top', 'w:top'),
                          ('right', 'w:right'), ('bottom', 'w:bottom')):
            if side not in paddings:
                continue
            size = self.bp._parse_size(paddings[side])
            if size is None:
                continue
            el = OxmlElement(tag)
            el.set(qn('w:w'), str(int(size.twips)))
            el.set(qn('w:type'), 'dxa')
            mar.append(el)
        if len(mar):
            tcPr.append(mar)

    def _apply_row_formats(self, table, rows_data):
        for r_idx, row in enumerate(rows_data):
            params = row['params']
            h = params.get('height')
            if h:
                size = self.bp._parse_size(h)
                if size is not None:
                    table.rows[r_idx].height = size
            if params.get('header', '').lower() in ('true', '1', 'yes'):
                trPr = table.rows[r_idx]._tr.get_or_add_trPr()
                for old in trPr.findall(qn('w:tblHeader')):
                    trPr.remove(old)
                th = OxmlElement('w:tblHeader')
                th.set(qn('w:val'), 'true')
                trPr.append(th)

    def _apply_table_format(self, table, table_params):
        spec = FormatSpec.parse(table_params.get('format', ''))

        if 'align' in spec:
            mapping = {
                'left': WD_TABLE_ALIGNMENT.LEFT,
                'center': WD_TABLE_ALIGNMENT.CENTER,
                'right': WD_TABLE_ALIGNMENT.RIGHT,
            }
            a = mapping.get(str(spec.arg('align')).lower())
            if a is not None:
                table.alignment = a

        if 'width' in spec:
            size = self.bp._parse_size(spec.arg('width'))
            if size is not None:
                table.autofit = False
                for row in table.rows:
                    for cell in row.cells:
                        cell.width = size

        b = spec.border()
        if b is None:
            # Дефолт: тонкая серая сетка, чтобы таблица не сливалась со страницей
            self._set_table_borders(table, '1pt', 'single', '#999999')
        else:
            self._set_table_borders(table, *b)

    def _set_table_borders(self, table, size_str='1pt',
                           style='single', color='#000000'):
        """
        size_str — '1pt', '0.5pt', '2' и т.п.
        style    — CSS или OOXML-имя ('solid' → 'single')
        color    — '#RRGGBB' или 'RRGGBB'
        """
        # 1/8 pt — единица измерения толщины линий в OOXML
        try:
            pt = float(re.sub(r'[^\d.]', '', str(size_str)) or '1')
        except ValueError:
            pt = 1.0
        sz = str(int(pt * 8))

        color = str(color).lstrip('#').upper()

        ooxml_style = self._BORDER_STYLE_MAP.get(
            str(style).lower(), 'single'
        )

        tblPr = table._tbl.tblPr
        for old in tblPr.findall(qn('w:tblBorders')):
            tblPr.remove(old)

        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right',
                     'insideH', 'insideV'):
            el = OxmlElement(f'w:{edge}')
            el.set(qn('w:val'), ooxml_style)
            el.set(qn('w:sz'), sz)
            el.set(qn('w:space'), '0')
            el.set(qn('w:color'), color)
            borders.append(el)

        # Правильная позиция в OOXML-схеме для CT_TblPr:
        # tblBorders идёт после tblW/jc/tblInd и ДО shd/tblLayout/tblCellMar/tblLook.
        successors = (
            'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
            'w:tblCaption', 'w:tblDescription', 'w:tblPrChange',
        )
        tblPr.insert_element_before(borders, *successors)

    def _set_table_borders(self, table, size_str='1pt', style='single', color='#000000'):
        """
        size_str — '1pt', '0.5pt', '2' и т.п.
        style    — 'single', 'dashed', 'dotted', 'double', 'none', ...
        color    — '#RRGGBB' или 'RRGGBB'
        """
        # 1/8 pt — единица измерения толщины линий в OOXML
        try:
            pt = float(re.sub(r'[^\d.]', '', str(size_str)) or '1')
        except ValueError:
            pt = 1.0
        sz = str(int(pt * 8))

        color = str(color).lstrip('#')

        tblPr = table._tbl.tblPr
        for old in tblPr.findall(qn('w:tblBorders')):
            tblPr.remove(old)

        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right',
                     'insideH', 'insideV'):
            el = OxmlElement(f'w:{edge}')
            el.set(qn('w:val'), style)
            el.set(qn('w:sz'), sz)
            el.set(qn('w:space'), '0')
            el.set(qn('w:color'), color)
            borders.append(el)
        tblPr.append(borders)


# ======================================================================
#                             СЕРИАЛИЗАТОР
# ======================================================================
class DocxTableSerializer:
    """Сериализует docx.Table в {block type:"table"}...{/block}."""

    def __init__(self, block_serializer):
        self.bs = block_serializer

    def serialize_table(self, table):
        """
        docx.Table -> строка вида:
            {block type:"table" format:"..."}
            {row ...}{cell ...}...{/cell}...{/row}
            ...
            {/block}
        """
        num_rows = len(table.rows)

        # 1. Собрать уникальные ячейки с их позициями и объединениями
        cell_map = {}
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                key = id(cell)
                if key not in cell_map:
                    cell_map[key] = {
                        'cell': cell,
                        'row': r, 'col': c,
                        'rowspan': 1, 'colspan': 1,
                    }
                else:
                    info = cell_map[key]
                    info['rowspan'] = max(
                        info['rowspan'], r - info['row'] + 1)
                    info['colspan'] = max(
                        info['colspan'], c - info['col'] + 1)

        # 2. Сгруппировать по строкам, отсортировать по колонкам
        per_row = [[] for _ in range(num_rows)]
        for info in cell_map.values():
            per_row[info['row']].append(info)
        for lst in per_row:
            lst.sort(key=lambda x: x['col'])

        # 3. Собрать строки
        row_strs = []
        for r in range(num_rows):
            row = table.rows[r]
            row_params = self._extract_row_params(row)
            cell_strs = []

            for info in per_row[r]:
                top, cell_spec = self._extract_cell_params(
                    info['cell'], info['colspan'], info['rowspan'])
                cell_body = self._serialize_cell_body(info['cell'])

                attrs = []
                top_str = self._dict_to_params(top)
                if top_str:
                    attrs.append(top_str)
                if cell_spec:
                    attrs.append(f'format:"{cell_spec.to_string()}"')

                if attrs:
                    cell_strs.append(
                        f'{{cell {" ".join(attrs)}}}'
                        f'{cell_body}{{/cell}}')
                else:
                    cell_strs.append(f'{{cell}}{cell_body}{{/cell}}')

            row_params_str = self._dict_to_params(row_params)
            inner = ''.join(cell_strs)
            if row_params_str:
                row_strs.append(
                    f'{{row {row_params_str}}}{inner}{{/row}}')
            else:
                row_strs.append(f'{{row}}{inner}{{/row}}')

        # 4. Шапка таблицы
        table_format = self._extract_table_format(table)
        params = ['type:"table"']
        if table_format:
            params.append(f'format:"{table_format}"')
        params_str = ' '.join(params)

        return (f'{{block {params_str}}}' + '\n' +
                '\n'.join(row_strs) + '\n{/block}')

    def _serialize_cell_body(self, cell):
        parts = []
        for child in cell._tc:
            if child.tag == qn('w:p'):
                p = Paragraph(child, cell)
                if not p.text.strip() and not p.runs:
                    continue
                parts.append(self.bs._serialize_paragraph(p))
            elif child.tag == qn('w:tbl'):
                t = DocxTable(child, cell)
                parts.append(self.serialize_table(t))
        return ''.join(parts)

    def _extract_row_params(self, row):
        params = {}
        trPr = row._tr.find(qn('w:trPr'))
        if trPr is None:
            return params
        h = trPr.find(qn('w:trHeight'))
        if h is not None:
            val = h.get(qn('w:val'))
            if val:
                try:
                    params['height'] = f'{int(val) / 20}pt'
                except ValueError:
                    pass
        if trPr.find(qn('w:tblHeader')) is not None:
            params['header'] = 'true'
        return params

    def _extract_cell_params(self, cell, colspan, rowspan):
        """
        Возвращает (top_params, FormatSpec).
        top_params — только colspan/rowspan.
        FormatSpec — всё остальное.
        """
        top = {}
        if colspan > 1:
            top['colspan'] = str(colspan)
        if rowspan > 1:
            top['rowspan'] = str(rowspan)

        spec = FormatSpec()
        tcPr = cell._tc.find(qn('w:tcPr'))
        if tcPr is None:
            return top, spec

        shd = tcPr.find(qn('w:shd'))
        if shd is not None:
            fill = shd.get(qn('w:fill'))
            if fill and fill.lower() not in ('auto', 'ffffff'):
                spec.set('bg-color', f'#{fill}')

        va = tcPr.find(qn('w:vAlign'))
        if va is not None:
            v = va.get(qn('w:val'))
            if v == 'center':
                spec.set('valign', 'middle')
            elif v in ('top', 'bottom'):
                spec.set('valign', v)

        jc = tcPr.find(qn('w:jc'))
        if jc is not None:
            v = jc.get(qn('w:val'))
            mapping = {'left': 'left', 'center': 'center',
                       'right': 'right', 'both': 'justify'}
            if v in mapping:
                spec.set('align', mapping[v])

        tcW = tcPr.find(qn('w:tcW'))
        if tcW is not None:
            w = tcW.get(qn('w:w'))
            typ = tcW.get(qn('w:type'))
            if w and typ == 'dxa':
                try:
                    spec.set('width', f'{round(int(w) / 20, 2)}pt')
                except ValueError:
                    pass

        tcMar = tcPr.find(qn('w:tcMar'))
        if tcMar is not None:
            for side, tag in (('left', 'w:left'), ('top', 'w:top'),
                              ('right', 'w:right'), ('bottom', 'w:bottom')):
                el = tcMar.find(qn(tag))
                if el is None:
                    continue
                w = el.get(qn('w:w'))
                typ = el.get(qn('w:type'))
                if w and typ == 'dxa':
                    try:
                        spec.set(f'padding-{side}',
                                 f'{round(int(w) / 20, 2)}pt')
                    except ValueError:
                        pass

        return top, spec

    def _extract_table_format(self, table) -> str:
        spec = FormatSpec()
        tblPr = table._tbl.find(qn('w:tblPr'))
        if tblPr is None:
            return ''

        jc = tblPr.find(qn('w:jc'))
        if jc is not None:
            v = jc.get(qn('w:val'))
            if v in ('left', 'center', 'right'):
                spec.set('align', v)

        tblW = tblPr.find(qn('w:tblW'))
        if tblW is not None:
            w = tblW.get(qn('w:w'))
            typ = tblW.get(qn('w:type'))
            if w and typ == 'dxa':
                try:
                    spec.set('width', f'{round(int(w) / 20, 2)}pt')
                except ValueError:
                    pass

        if tblPr.find(qn('w:tblBorders')) is not None:
            spec.set('border', '1pt,single,#000000')

        return spec.to_string()

    def _dict_to_params(self, d):
        return ' '.join(f'{k}:"{v}"' for k, v in d.items())